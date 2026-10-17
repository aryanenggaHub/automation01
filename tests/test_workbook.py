from datetime import datetime
from pathlib import Path

import openpyxl

from splunk_seq.config import OutputCfg
from splunk_seq.models import AMBIGUOUS, ERROR, NO_SESSION, OK, ApiResult, SessionResult
from splunk_seq.workbook import SequenceWorkbook, output_path, sheet_name


def row(ts, gw="", **extra):
    return {"Timestamp": ts, "Gateway API Endpoint": gw, **extra}


def two_session_result():
    return ApiResult(
        input_api="/v1/a",
        status=OK,
        sessions=[
            SessionResult("s1", [row("01.000", "A"), row("01.100", "B")], [AMBIGUOUS, ""]),
            SessionResult("s2", [row("02.000", "C")], [""]),
        ],
    )


def build(tmp_path, *results, total=None):
    path = tmp_path / "out" / "wb.xlsx"
    wb = SequenceWorkbook(path, total or len(results))
    for i, result in enumerate(results, start=1):
        wb.add(i, result)
        wb.save()
    return openpyxl.load_workbook(path)


def test_sheet_name_rules():
    assert sheet_name(1, 5, "/v1/customers/{id}/accounts") == "001_v1_customers_id_accounts"
    assert len(sheet_name(1, 5, "/v1/" + "x" * 60)) == 31
    assert sheet_name(12, 1500, "/a") == "0012_a"
    assert sheet_name(3, 5, "/") == "003"


def test_output_path():
    assert output_path(OutputCfg("output", "api_sequence"), datetime(2026, 10, 17, 14, 23, 1)) == Path(
        "output/api_sequence_20261017_142301.xlsx"
    )


def test_index_sheet_first_with_header_counts_and_hyperlink(tmp_path):
    wb = build(tmp_path, two_session_result())
    ws = wb["Index"]
    assert wb.sheetnames[0] == "Index"
    assert [c.value for c in ws[1]] == ["#", "Input API", "Status", "Sessions found", "Rows", "Ambiguous rows", "Error"]
    assert [c.value for c in ws[2]] == [1, "/v1/a", "OK", 2, 3, 1, None]
    assert ws["B2"].hyperlink.location == "'001_v1_a'!A1"


def test_api_sheet_layout(tmp_path):
    ws = build(tmp_path, two_session_result())["001_v1_a"]
    assert ws["A1"].value == "/v1/a"
    assert ws["A2"].hyperlink.location == "'Index'!A1"
    assert [c.value for c in ws[4]] == ["Session ID", "Timestamp", "Gateway API Endpoint", "Flag"]
    assert [c.value for c in ws[5]] == ["s1", "01.000", "A", "AMBIGUOUS"]
    assert [c.value for c in ws[6]] == ["s1", "01.100", "B", None]
    assert all(c.value is None for c in ws[7])
    assert [c.value for c in ws[8]] == ["s2", "02.000", "C", None]
    assert ws.freeze_panes == "A5"
    assert ws.auto_filter.ref == "A4:D8"


def test_column_union_appends_new_columns_right(tmp_path):
    result = ApiResult("/v1/a", OK, [
        SessionResult("s1", [row("01", "A")], [""]),
        SessionResult("s2", [row("02", "B", X="x")], [""]),
    ])
    ws = build(tmp_path, result)["001_v1_a"]
    assert [c.value for c in ws[4]] == ["Session ID", "Timestamp", "Gateway API Endpoint", "X", "Flag"]
    assert [c.value for c in ws[7]] == ["s2", "02", "B", "x", None]


def test_no_session_still_gets_sheet(tmp_path):
    wb = build(tmp_path, ApiResult("/v1/none", NO_SESSION))
    assert wb["001_v1_none"]["A1"].value == "/v1/none"
    assert [c.value for c in wb["Index"][2]] == [1, "/v1/none", "NO_SESSION", 0, 0, 0, None]


def test_error_message_in_index(tmp_path):
    wb = build(tmp_path, ApiResult("/v1/a", ERROR, error="HTTP 500"))
    assert wb["Index"]["C2"].value == "ERROR" and wb["Index"]["G2"].value == "HTTP 500"


def test_formula_like_values_stay_text(tmp_path):
    result = ApiResult("/v1/a", OK, [SessionResult("s1", [row("01", "=HYPERLINK(1)")], [""])])
    cell = build(tmp_path, result)["001_v1_a"]["C5"]
    assert cell.value == "=HYPERLINK(1)" and cell.data_type == "s"


def test_illegal_control_chars_stripped(tmp_path):
    result = ApiResult("/v1/a", OK, [SessionResult("s1", [row("01", "ab\x01c")], [""])])
    assert build(tmp_path, result)["001_v1_a"]["C5"].value == "abc"


def test_save_after_each_add_keeps_earlier_sheets(tmp_path):
    wb = build(tmp_path, two_session_result(), ApiResult("/v1/b", NO_SESSION))
    assert wb.sheetnames == ["Index", "001_v1_a", "002_v1_b"]
    assert wb["Index"]["B3"].value == "/v1/b"
