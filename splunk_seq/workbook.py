"""Write the output workbook: an Index sheet plus one sheet per Input API (spec §8)."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.hyperlink import Hyperlink

from .config import OutputCfg
from .models import AMBIGUOUS, ApiResult

INDEX_SHEET = "Index"
INDEX_HEADER = ["#", "Input API", "Status", "Sessions found", "Rows", "Ambiguous rows", "Error"]
HEADER_ROW = 4
MAX_SHEET_NAME = 31


def sheet_name(index: int, total: int, endpoint: str) -> str:
    prefix = f"{index:0{max(3, len(str(total)))}d}"
    cleaned = re.sub(r"[^A-Za-z0-9]+", "_", endpoint).strip("_")
    name = f"{prefix}_{cleaned}" if cleaned else prefix
    return name[:MAX_SHEET_NAME]


def output_path(cfg: OutputCfg, now: datetime) -> Path:
    return Path(cfg.dir) / f"{cfg.filename_prefix}_{now:%Y%m%d_%H%M%S}.xlsx"


def _set(ws, row: int, col: int, value):
    """Write a value as literal data: strip characters Excel rejects, never treat text as a formula."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell.data_type = "s"
    else:
        cell.value = value
    return cell


def _link(cell, sheet: str) -> None:
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!A1")
    cell.style = "Hyperlink"


class SequenceWorkbook:
    def __init__(self, path: Path, total: int):
        self.path = Path(path)
        self.total = total
        self._wb = openpyxl.Workbook()
        index = self._wb.active
        index.title = INDEX_SHEET
        for col, title in enumerate(INDEX_HEADER, start=1):
            _set(index, 1, col, title)
        index.freeze_panes = "A2"

    def add(self, index: int, result: ApiResult) -> None:
        name = sheet_name(index, self.total, result.input_api)
        ws = self._wb.create_sheet(name)
        _set(ws, 1, 1, result.input_api)
        _link(_set(ws, 2, 1, "← Index"), INDEX_SHEET)

        columns: list[str] = []
        for session in result.sessions:
            for r in session.rows:
                columns.extend(c for c in r if c not in columns)
        header = ["Session ID", *columns, "Flag"]
        for col, title in enumerate(header, start=1):
            _set(ws, HEADER_ROW, col, title)

        row_no = HEADER_ROW + 1
        for n, session in enumerate(result.sessions):
            if n:
                row_no += 1  # blank row between Sessions
            for r, flag in zip(session.rows, session.flags):
                values = [session.session_id, *(r.get(c, "") for c in columns), flag]
                for col, value in enumerate(values, start=1):
                    if value != "":
                        _set(ws, row_no, col, value)
                row_no += 1
        last_row = max(row_no - 1, HEADER_ROW)
        ws.freeze_panes = f"A{HEADER_ROW + 1}"
        ws.auto_filter.ref = f"A{HEADER_ROW}:{ws.cell(row=HEADER_ROW, column=len(header)).column_letter}{last_row}"

        index_ws = self._wb[INDEX_SHEET]
        rows = sum(len(s.rows) for s in result.sessions)
        ambiguous = sum(f == AMBIGUOUS for s in result.sessions for f in s.flags)
        values = [index, result.input_api, result.status, len(result.sessions), rows, ambiguous, result.error or None]
        for col, value in enumerate(values, start=1):
            _set(index_ws, index + 1, col, value)
        _link(index_ws.cell(row=index + 1, column=2), name)

    def save(self) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._wb.save(self.path)
